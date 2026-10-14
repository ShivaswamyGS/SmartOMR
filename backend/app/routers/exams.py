import csv
import io
from fastapi import APIRouter, Depends, HTTPException, status, UploadFile, File
from sqlalchemy.orm import Session
from typing import List, Optional

from ..database import get_db
from ..models import Exam, AnswerKey, RosterEntry, User, EvaluationResult, AuditLog
from ..schemas import ExamCreate, ExamUpdate, ExamResponse, AnswerKeyCreate, AnswerKeyResponse, RosterItem
from ..security import get_current_user

router = APIRouter(prefix="/exams", tags=["Exams"])

@router.post("", response_model=ExamResponse)
def create_exam(
    exam_in: ExamCreate, 
    db: Session = Depends(get_db), 
    current_user: User = Depends(get_current_user)
):
    exam = Exam(
        user_id=current_user.id,
        title=exam_in.title,
        subject=exam_in.subject,
        code=exam_in.code,
        total_questions=exam_in.total_questions,
        marks_per_correct=exam_in.marks_per_correct,
        negative_marks_per_incorrect=exam_in.negative_marks_per_incorrect,
        max_marks=exam_in.max_marks,
        allow_partial_marking=exam_in.allow_partial_marking,
        template_id=exam_in.template_id,
        status="active"
    )
    db.add(exam)
    db.commit()
    db.refresh(exam)
    
    # Audit Log
    db.add(AuditLog(
        user_id=current_user.id,
        action="CREATE_EXAM",
        entity_type="Exam",
        entity_id=str(exam.id),
        details=f"Created exam '{exam.title}' with {exam.total_questions} questions"
    ))
    db.commit()
    
    return exam

@router.get("", response_model=List[ExamResponse])
def list_exams(
    db: Session = Depends(get_db), 
    current_user: User = Depends(get_current_user)
):
    query = db.query(Exam)
    if current_user.role != "admin":
        query = query.filter(Exam.user_id == current_user.id)
    exams = query.order_by(Exam.created_at.desc()).all()
    
    # Enrich with computed properties
    res = []
    for ex in exams:
        sub_count = db.query(EvaluationResult).filter(EvaluationResult.exam_id == ex.id).count()
        results = db.query(EvaluationResult).filter(EvaluationResult.exam_id == ex.id).all()
        avg = round(sum(r.raw_score for r in results) / len(results), 2) if results else 0.0
        
        dto = ExamResponse(
            id=ex.id,
            user_id=ex.user_id,
            title=ex.title,
            subject=ex.subject,
            code=ex.code,
            total_questions=ex.total_questions,
            marks_per_correct=ex.marks_per_correct,
            negative_marks_per_incorrect=ex.negative_marks_per_incorrect,
            max_marks=ex.max_marks,
            allow_partial_marking=ex.allow_partial_marking,
            template_id=ex.template_id,
            status=ex.status,
            created_at=ex.created_at,
            updated_at=ex.updated_at,
            has_answer_key=ex.answer_key is not None,
            total_submissions=sub_count,
            average_score=avg
        )
        res.append(dto)
    return res

@router.get("/{exam_id}", response_model=ExamResponse)
def get_exam(
    exam_id: int, 
    db: Session = Depends(get_db), 
    current_user: User = Depends(get_current_user)
):
    exam = db.query(Exam).filter(Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")
        
    sub_count = db.query(EvaluationResult).filter(EvaluationResult.exam_id == exam.id).count()
    results = db.query(EvaluationResult).filter(EvaluationResult.exam_id == exam.id).all()
    avg = round(sum(r.raw_score for r in results) / len(results), 2) if results else 0.0
    
    dto = ExamResponse(
        id=exam.id,
        user_id=exam.user_id,
        title=exam.title,
        subject=exam.subject,
        code=exam.code,
        total_questions=exam.total_questions,
        marks_per_correct=exam.marks_per_correct,
        negative_marks_per_incorrect=exam.negative_marks_per_incorrect,
        max_marks=exam.max_marks,
        allow_partial_marking=exam.allow_partial_marking,
        template_id=exam.template_id,
        status=exam.status,
        created_at=exam.created_at,
        updated_at=exam.updated_at,
        has_answer_key=exam.answer_key is not None,
        total_submissions=sub_count,
        average_score=avg
    )
    return dto

@router.put("/{exam_id}", response_model=ExamResponse)
def update_exam(
    exam_id: int,
    exam_in: ExamUpdate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    exam = db.query(Exam).filter(Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")
        
    for field, val in exam_in.dict(exclude_unset=True).items():
        setattr(exam, field, val)
        
    db.commit()
    db.refresh(exam)
    return exam

# --- Answer Key Management ---
@router.post("/{exam_id}/answer-key", response_model=AnswerKeyResponse)
def save_answer_key(
    exam_id: int,
    key_in: AnswerKeyCreate,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    exam = db.query(Exam).filter(Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")
        
    # Validate that answers are within valid choices
    valid_options = {"A", "B", "C", "D", "E"}
    for q_str, choice in key_in.answers.items():
        if choice and choice.upper() not in valid_options:
            raise HTTPException(
                status_code=400,
                detail=f"Invalid option '{choice}' for question {q_str}. Expected one of {valid_options}."
            )
            
    existing_key = db.query(AnswerKey).filter(AnswerKey.exam_id == exam_id).first()
    if existing_key:
        existing_key.total_questions = key_in.total_questions
        existing_key.answers = key_in.answers
        existing_key.custom_weights = key_in.custom_weights
        db.commit()
        db.refresh(existing_key)
        return existing_key
    else:
        new_key = AnswerKey(
            exam_id=exam_id,
            total_questions=key_in.total_questions,
            answers=key_in.answers,
            custom_weights=key_in.custom_weights
        )
        db.add(new_key)
        db.commit()
        db.refresh(new_key)
        return new_key

@router.get("/{exam_id}/answer-key", response_model=AnswerKeyResponse)
def get_answer_key(
    exam_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    key = db.query(AnswerKey).filter(AnswerKey.exam_id == exam_id).first()
    if not key:
        raise HTTPException(status_code=404, detail="Answer key not configured for this exam")
    return key

# --- Roster Upload Management ---
@router.post("/{exam_id}/roster")
async def upload_roster_file(
    exam_id: int,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    exam = db.query(Exam).filter(Exam.id == exam_id).first()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")
        
    contents = await file.read()
    filename = file.filename.lower()
    
    entries = []
    
    if filename.endswith(".csv"):
        text = contents.decode("utf-8-sig", errors="ignore")
        reader = csv.DictReader(io.StringIO(text))
        for row in reader:
            # Flexible field matching
            sid = row.get("Student ID") or row.get("student_id") or row.get("Roll Number") or row.get("roll_number") or row.get("ID")
            sname = row.get("Student Name") or row.get("student_name") or row.get("Name") or row.get("name")
            roll = row.get("Roll Number") or row.get("roll_number") or sid
            email = row.get("Email") or row.get("email")
            sec = row.get("Class") or row.get("Section") or row.get("class_section")
            if sid and sname:
                entries.append(RosterEntry(
                    exam_id=exam_id,
                    student_id=str(sid).strip(),
                    student_name=str(sname).strip(),
                    roll_number=str(roll).strip() if roll else None,
                    email=str(email).strip() if email else None,
                    class_section=str(sec).strip() if sec else None
                ))
    elif filename.endswith(".xlsx"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
        ws = wb.active
        headers = [str(cell.value or '').strip().lower() for cell in next(ws.iter_rows(max_row=1))]
        
        sid_col = next((i for i, h in enumerate(headers) if "id" in h or "roll" in h), None)
        sname_col = next((i for i, h in enumerate(headers) if "name" in h), None)
        
        if sid_col is not None and sname_col is not None:
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row[sid_col] and row[sname_col]:
                    entries.append(RosterEntry(
                        exam_id=exam_id,
                        student_id=str(row[sid_col]).strip(),
                        student_name=str(row[sname_col]).strip(),
                        roll_number=str(row[sid_col]).strip()
                    ))
    else:
        raise HTTPException(status_code=400, detail="Only CSV or XLSX rosters are supported.")
        
    if not entries:
        raise HTTPException(status_code=400, detail="No valid student rows found in uploaded roster.")
        
    # Remove existing roster entries for exam and replace
    db.query(RosterEntry).filter(RosterEntry.exam_id == exam_id).delete()
    for e in entries:
        db.add(e)
    db.commit()
    
    return {"message": f"Successfully imported {len(entries)} students into roster", "count": len(entries)}

@router.get("/{exam_id}/roster")
def get_roster(
    exam_id: int,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user)
):
    roster = db.query(RosterEntry).filter(RosterEntry.exam_id == exam_id).all()
    return [{"student_id": r.student_id, "student_name": r.student_name, "roll_number": r.roll_number, "email": r.email, "class_section": r.class_section} for r in roster]
